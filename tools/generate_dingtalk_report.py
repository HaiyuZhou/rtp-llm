#!/usr/bin/env python3
"""Build a DingTalk-friendly DOCX/PDF bundle from a performance Markdown report."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report-dir", type=Path, required=True)
    parser.add_argument("--source-json", type=Path, required=True)
    parser.add_argument("--font", type=Path, required=True)
    parser.add_argument("--expected-runs", type=int, default=None)
    parser.add_argument("--model-label", default="Model")
    return parser.parse_args()


def add_chart_previews(
    report_dir: Path,
    source_json: Path,
    expected_runs: int | None,
    model_label: str = "Model",
) -> dict[str, Path]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    from rtp_llm.test.perf_test.cache_grid.plot.generate_prefill_interactive_chart import (
        load_rows,
    )

    rows = load_rows(source_json, batch_size=1, all_runs=True)
    if expected_runs is not None and len(rows) != expected_runs:
        raise ValueError(f"expected {expected_runs} successful runs, found {len(rows)}")

    compute = np.asarray([row["compute_len"] for row in rows], dtype=float)
    cache = np.asarray([row["cache_len"] for row in rows], dtype=float)
    input_tokens = compute + cache
    rt = np.asarray([row["prefill_rt"] for row in rows], dtype=float)
    metrics = {
        "rt": (
            rt,
            "Engine prefill latency (ms)",
            "08_prefill_rt_3d_all_runs_preview.png",
        ),
        "compute": (
            compute / rt * 60000.0,
            "Whole-system uncached-compute TPM",
            "09_compute_tpm_3d_all_runs_preview.png",
        ),
        "effective": (
            input_tokens / rt * 60000.0,
            "Whole-system effective-input TPM",
            "10_effective_tpm_3d_all_runs_preview.png",
        ),
    }
    outputs: dict[str, Path] = {}
    for key, (values, label, filename) in metrics.items():
        fig = plt.figure(figsize=(11.5, 7.5), constrained_layout=True)
        axis = fig.add_subplot(111, projection="3d")
        colors = np.log10(np.maximum(values, 1e-12))
        scatter = axis.scatter(
            compute / 1024.0,
            cache / 1024.0,
            values,
            c=colors,
            cmap="viridis",
            s=0.35,
            alpha=0.35,
            linewidths=0,
            rasterized=True,
        )
        axis.set_xlabel("Uncached compute tokens (K)")
        axis.set_ylabel("Observed cached tokens (K)")
        axis.set_zlabel(label)
        axis.set_title(
            f"{model_label} Prefill — {label} — all {len(rows):,} formal runs"
        )
        axis.view_init(elev=24, azim=-128)
        colorbar = fig.colorbar(scatter, ax=axis, shrink=0.62, pad=0.08)
        colorbar.set_label(f"log10({label})")
        output = report_dir / "charts" / filename
        fig.savefig(output, dpi=180, facecolor="white")
        plt.close(fig)
        outputs[key] = output
    return outputs


def add_fit_gap_preview(report_dir: Path) -> Path:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    prediction_path = report_dir / "formula_restricted_symbolic" / "predictions.csv"
    targets: list[float] = []
    predictions: list[float] = []
    apes: list[float] = []
    splits: list[str] = []
    with prediction_path.open(newline="", encoding="utf-8") as stream:
        for row in csv.DictReader(stream):
            targets.append(float(row["target_ms"]))
            predictions.append(float(row["predicted_ms"]))
            apes.append(float(row["ape_pct"]))
            splits.append(row["split"])

    target = np.asarray(targets)
    predicted = np.asarray(predictions)
    ape = np.asarray(apes)
    fig, axes = plt.subplots(1, 2, figsize=(12, 5.2), constrained_layout=True)
    palette = {"train": "#4c78a8", "validation": "#f58518", "test": "#54a24b"}
    for split in ("train", "validation", "test"):
        mask = np.asarray([value == split for value in splits])
        axes[0].scatter(
            target[mask],
            predicted[mask],
            s=2,
            alpha=0.25,
            linewidths=0,
            label=f"{split} ({int(mask.sum()):,})",
            color=palette[split],
            rasterized=True,
        )
    maximum = float(max(target.max(), predicted.max()))
    axes[0].plot([0, maximum], [0, maximum], color="black", linewidth=1, linestyle="--")
    axes[0].set_xlabel("Measured client wall TTFT median (ms)")
    axes[0].set_ylabel("Predicted client wall TTFT (ms)")
    axes[0].set_title("Restricted symbolic fit: measured vs predicted")
    axes[0].legend(markerscale=4)
    axes[0].grid(alpha=0.2)

    capped = np.minimum(ape, 20.0)
    axes[1].hist(capped, bins=80, color="#4c78a8", alpha=0.85)
    axes[1].axvline(10.0, color="#e45756", linestyle="--", label="10% APE")
    axes[1].set_xlabel("Absolute percentage error (%) — values >20% clipped")
    axes[1].set_ylabel("Geometry count")
    axes[1].set_title(f"APE distribution, all {len(ape):,} geometries")
    axes[1].set_yscale("log")
    axes[1].legend()
    axes[1].grid(alpha=0.2)

    output = report_dir / "formula_restricted_symbolic" / "fit_gap.png"
    fig.savefig(output, dpi=180, facecolor="white")
    plt.close(fig)
    return output


def clean_inline(text: str) -> str:
    return text.replace("**", "").replace("`", "")


def markdown_elements(
    markdown: str, report_dir: Path, generated: dict[str, Path]
) -> list[tuple[Any, ...]]:
    image_map = {
        "### 全局分布": [
            (
                report_dir / "charts/01_latency_landscape.png",
                "图 1：Engine/client latency landscape",
            ),
            (report_dir / "charts/02_cold_scaling.png", "图 2：Cold request scaling"),
        ],
        "### Cache 命中对比": [
            (
                report_dir / "charts/03_cache_speedup.png",
                "图 3：同 input length 的 cache speedup",
            ),
            (
                report_dir / "charts/06_compute_scaling.png",
                "图 4：按 cache ratio 分层的 compute scaling",
            ),
        ],
        "### 重复性": [
            (
                report_dir / "charts/04_timing_scope_comparison.png",
                "图 5：Client wall 与 engine prefill 口径",
            ),
            (report_dir / "charts/05_repeatability.png", "图 6：三轮正式测量的重复性"),
        ],
        "## 受限符号拟合候选": [
            (generated["fit"], "图 7：受限符号回归的预测误差"),
        ],
        "## 图表清单": [
            (
                generated["rt"],
                "图 8：全量 engine prefill latency 静态预览；交互版本见 HTML 附件",
            ),
            (
                generated["compute"],
                "图 9：全量 whole-system compute TPM 静态预览；交互版本见 HTML 附件",
            ),
            (
                generated["effective"],
                "图 10：全量 whole-system effective-input TPM 静态预览；交互版本见 HTML 附件",
            ),
        ],
    }
    lines = markdown.splitlines()
    elements: list[tuple[Any, ...]] = []
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if (
            line.startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\|[|:\- ]+\|$", lines[index + 1])
        ):
            rows = [[clean_inline(cell.strip()) for cell in line.strip("|").split("|")]]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(
                    [
                        clean_inline(cell.strip())
                        for cell in lines[index].strip("|").split("|")
                    ]
                )
                index += 1
            elements.append(("table", rows))
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            elements.append(("heading", level, clean_inline(heading.group(2))))
            for path, caption in image_map.get(line, []):
                elements.append(("image", path, caption))
            index += 1
            continue
        bullet = re.match(r"^-\s+(.*)$", line)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if bullet:
            elements.append(("bullet", bullet.group(1)))
        elif numbered:
            elements.append(("number", numbered.group(1), numbered.group(2)))
        else:
            elements.append(("paragraph", line))
        index += 1
    return elements


def add_inline_runs(paragraph: Any, text: str) -> None:
    chunks = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for chunk in chunks:
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = paragraph.style.font.size
        else:
            paragraph.add_run(chunk)


def build_docx(
    report_dir: Path, elements: Iterable[tuple[Any, ...]], model_label: str = "Model"
) -> Path:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    for style_name, size, color in (
        ("Normal", 10.5, "222222"),
        ("Title", 22, "17365D"),
        ("Heading 1", 17, "17365D"),
        ("Heading 2", 14, "2F5597"),
        ("Heading 3", 12, "4472C4"),
        ("Caption", 9, "666666"),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for kind, *payload in elements:
        if kind == "heading":
            level, text = payload
            paragraph = document.add_paragraph(
                style="Title" if level == 1 else f"Heading {min(level - 1, 3)}"
            )
            paragraph.paragraph_format.keep_with_next = True
            add_inline_runs(paragraph, text)
        elif kind == "paragraph":
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline_runs(paragraph, payload[0])
        elif kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, payload[0])
        elif kind == "number":
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, payload[1])
        elif kind == "table":
            rows = payload[0]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.style = document.styles["Normal"]
                        for run in paragraph.runs:
                            run.font.size = Pt(8.5)
                            if row_index == 0:
                                run.bold = True
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "D9EAF7")
                        cell._tc.get_or_add_tcPr().append(shading)
            document.add_paragraph()
        elif kind == "image":
            path, caption = payload
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Cm(16.7))
            caption_paragraph = document.add_paragraph(caption, style="Caption")
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_paragraph.paragraph_format.keep_with_next = False

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{model_label} Prefill 全量测试分析 · DingTalk delivery copy")
    output = report_dir / "REPORT_DINGTALK.docx"
    document.core_properties.title = f"{model_label} Prefill 全量测试分析"
    document.core_properties.subject = "DingTalk-ready report with embedded charts"
    document.save(output)
    return output


def wrap_text(draw: Any, text: str, font: Any, width: int) -> list[str]:
    tokens = re.findall(r"[\u3400-\u9fff]|\s+|[^\u3400-\u9fff\s]+", clean_inline(text))
    lines: list[str] = []
    current = ""
    for token in tokens:
        candidate = current + token
        if current and draw.textlength(candidate, font=font) > width:
            lines.append(current.rstrip())
            current = token.lstrip()
        else:
            current = candidate
    if current:
        lines.append(current.rstrip())
    return lines or [""]


def build_pdf(
    report_dir: Path, elements: Iterable[tuple[Any, ...]], font_path: Path
) -> Path:
    from PIL import Image, ImageDraw, ImageFont

    page_width, page_height = 1240, 1754
    margin, content_width = 82, 1076
    fonts = {
        "title": ImageFont.truetype(str(font_path), 42),
        "h1": ImageFont.truetype(str(font_path), 32),
        "h2": ImageFont.truetype(str(font_path), 27),
        "h3": ImageFont.truetype(str(font_path), 23),
        "body": ImageFont.truetype(str(font_path), 20),
        "small": ImageFont.truetype(str(font_path), 16),
        "caption": ImageFont.truetype(str(font_path), 17),
    }
    pages: list[Image.Image] = []
    page = Image.new("RGB", (page_width, page_height), "white")
    draw = ImageDraw.Draw(page)
    y = margin

    def new_page() -> None:
        nonlocal page, draw, y
        pages.append(page)
        page = Image.new("RGB", (page_width, page_height), "white")
        draw = ImageDraw.Draw(page)
        y = margin

    def ensure(height: int) -> None:
        if y + height > page_height - margin:
            new_page()

    def text_block(
        text: str, font: Any, indent: int = 0, spacing: int = 8, fill: str = "#222222"
    ) -> None:
        nonlocal y
        lines = wrap_text(draw, text, font, content_width - indent)
        line_height = font.size + spacing
        ensure(line_height * len(lines) + spacing)
        for line in lines:
            draw.text((margin + indent, y), line, font=font, fill=fill)
            y += line_height
        y += spacing

    for kind, *payload in elements:
        if kind == "heading":
            level, text = payload
            font = fonts["title"] if level == 1 else fonts[f"h{min(level - 1, 3)}"]
            text_block(text, font, spacing=14, fill="#17365D")
        elif kind == "paragraph":
            text_block(payload[0], fonts["body"])
        elif kind == "bullet":
            text_block("• " + payload[0], fonts["body"], indent=18)
        elif kind == "number":
            text_block(f"{payload[0]}. {payload[1]}", fonts["body"], indent=18)
        elif kind == "image":
            path, caption = payload
            image = Image.open(path).convert("RGB")
            ratio = min(content_width / image.width, 760 / image.height, 1.0)
            size = (max(1, int(image.width * ratio)), max(1, int(image.height * ratio)))
            image = image.resize(size, Image.Resampling.LANCZOS)
            ensure(image.height + 70)
            page.paste(image, (margin + (content_width - image.width) // 2, y))
            y += image.height + 8
            caption_lines = wrap_text(draw, caption, fonts["caption"], content_width)
            for line in caption_lines:
                line_width = draw.textlength(line, font=fonts["caption"])
                draw.text(
                    ((page_width - line_width) / 2, y),
                    line,
                    font=fonts["caption"],
                    fill="#666666",
                )
                y += fonts["caption"].size + 5
            y += 15
        elif kind == "table":
            rows = payload[0]
            columns = len(rows[0])
            cell_width = content_width // columns
            font = fonts["small"]
            for row_index, row in enumerate(rows):
                wrapped = [
                    wrap_text(draw, value, font, cell_width - 12) for value in row
                ]
                row_height = max(len(lines) for lines in wrapped) * (font.size + 4) + 12
                ensure(row_height)
                for column_index, lines in enumerate(wrapped):
                    x = margin + column_index * cell_width
                    fill = "#D9EAF7" if row_index == 0 else "white"
                    draw.rectangle(
                        (x, y, x + cell_width, y + row_height),
                        fill=fill,
                        outline="#999999",
                        width=1,
                    )
                    line_y = y + 6
                    for line in lines:
                        draw.text((x + 6, line_y), line, font=font, fill="#222222")
                        line_y += font.size + 4
                y += row_height
            y += 14

    pages.append(page)
    output = report_dir / "REPORT_DINGTALK.pdf"
    pages[0].save(output, "PDF", resolution=150, save_all=True, append_images=pages[1:])
    return output


def update_manifest(
    report_dir: Path, generated: dict[str, Path], documents: list[Path]
) -> None:
    path = report_dir / "MANIFEST.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    charts = manifest.setdefault("charts", [])
    for key in ("rt", "compute", "effective"):
        relative = generated[key].relative_to(report_dir).as_posix()
        if relative not in charts:
            charts.append(relative)
    formula = manifest.setdefault("formula", [])
    fit_relative = generated["fit"].relative_to(report_dir).as_posix()
    if fit_relative not in formula:
        formula.append(fit_relative)
    manifest["documents"] = [
        document.relative_to(report_dir).as_posix() for document in documents
    ]
    path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def main() -> int:
    args = parse_args()
    report_dir = args.report_dir.resolve()
    previews = add_chart_previews(
        report_dir, args.source_json.resolve(), args.expected_runs, args.model_label
    )
    previews["fit"] = add_fit_gap_preview(report_dir)
    markdown = (report_dir / "REPORT.md").read_text(encoding="utf-8")
    elements = markdown_elements(markdown, report_dir, previews)
    docx = build_docx(report_dir, elements, args.model_label)
    pdf = build_pdf(report_dir, elements, args.font.resolve())
    update_manifest(report_dir, previews, [docx, pdf])
    result = {
        "docx": str(docx),
        "pdf": str(pdf),
        "previews": {key: str(value) for key, value in previews.items()},
    }
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
